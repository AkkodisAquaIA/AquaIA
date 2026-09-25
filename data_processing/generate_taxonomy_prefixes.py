#!/usr/bin/env python3

"""Génère le petit registre JSON utilisé pour décoder les classes CVAT.
Le json est déjà généré par le script `generate_taxonomy_prefixes.py` à partir du fichier Excel de référence.
Le fichier généré se trouve dans `data_processing/resources/taxonomy_prefixes.json` et est utilisé par `crop_detection_to_classification.py` pour décoder les classes CVAT."""

import argparse
import json
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


CLASS_COLUMN = "Classe"
CANONICAL_OVERRIDES = {
    (2, "Rhync"): "Rhynchobdellida",
    (2, "Union"): "Unionida",
}


def normalize_text(value):
    """Supprime les espaces inutiles et normalise les accents."""
    return unicodedata.normalize("NFC", str(value).strip())


def find_taxonomy_sheet(workbook):
    """Trouve la feuille et la colonne des noms taxonomiques complets."""
    for sheet in workbook.worksheets:
        headers = [normalize_text(cell.value) if cell.value is not None else "" for cell in sheet[1]]
        if CLASS_COLUMN in headers:
            return sheet, headers.index(CLASS_COLUMN)
    raise ValueError(f"Colonne absente du classeur : {CLASS_COLUMN}")


def build_prefix_registry(excel_path):
    """Construit une correspondance unique pour chaque rang taxonomique."""
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    sheet, class_column_index = find_taxonomy_sheet(workbook)
    candidates = [dict(), dict(), dict()]

    for row in sheet.iter_rows(min_row=2, values_only=True):
        value = row[class_column_index]
        if value is None:
            continue
        parts = normalize_text(value).split("_")
        if len(parts) < 3:
            raise ValueError(f"Nom taxonomique incomplet : {value}")

        for rank_index, full_name in enumerate(parts[:3]):
            prefix = full_name[:5]
            candidates[rank_index].setdefault(prefix, set()).add(full_name)

    workbook.close()
    prefixes = []

    for rank_index, rank_candidates in enumerate(candidates):
        rank_prefixes = {}
        for prefix, names in sorted(rank_candidates.items()):
            override = CANONICAL_OVERRIDES.get((rank_index, prefix))
            if override is not None:
                rank_prefixes[prefix] = override
            elif len(names) == 1:
                rank_prefixes[prefix] = next(iter(names))
            else:
                options = ", ".join(sorted(names))
                raise ValueError(f"Préfixe ambigu au rang {rank_index + 1} : {prefix} ({options})")
        prefixes.append(rank_prefixes)

    return {
        "version": 1,
        "prefix_length": 5,
        "ranks": ["embranchement", "classe", "ordre"],
        "prefixes": prefixes,
    }


def parse_args():
    parser = argparse.ArgumentParser(description="Génère le registre JSON des préfixes taxonomiques.")
    parser.add_argument("excel_path", type=Path, help="Classeur contenant la taxonomie de référence")
    parser.add_argument("output_path", type=Path, help="Fichier JSON à créer")
    return parser.parse_args()


def main():
    args = parse_args()
    registry = build_prefix_registry(args.excel_path.resolve())
    args.output_path.parent.mkdir(parents=True, exist_ok=True)
    with args.output_path.open("w", encoding="utf-8") as file:
        json.dump(registry, file, ensure_ascii=False, indent=2)
        file.write("\n")
    print(f"Registre créé : {args.output_path}")


if __name__ == "__main__":
    main()
