
from pathlib import Path

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font
from openpyxl.utils import column_index_from_string
from openpyxl.styles import Font, Alignment

# ============================================================
# PARAMETRES
# ============================================================

chemin = Path(
    "C:/Users/pierre.fancelli/Documents/_Dev/Aqua-IA/__Excel__/Comp_D"
)

fichier_akkodis = chemin / "Data_Akkodis.xlsx"
fichier_adour = chemin / "Data_Adour.xlsx"
fichier_sortie = chemin / "Ak_VS_Ad_Conplet__.xlsx"


# ============================================================
# FONCTIONS
# ============================================================

def nb_premieres_lettres(texte, nb):
    """
    Retourne les nb premières lettres du champ.
    """
    if texte is None:
        return ""

    texte = str(texte).strip()

    if not texte:
        return ""

    return texte[:nb]


def champ_adour(texte, remplacement):
    """
    Traite un champ provenant du fichier Adour.

    Si le champ commence par 'Sans ', il est remplacé
    par la valeur indiquée dans 'remplacement'.
    """
    if texte is None:
        return remplacement

    texte = str(texte).strip()

    if texte.lower().startswith("sans"):
        return remplacement

    return texte


def construire_taxon_akkodis(taxon):
    """
    Transforme :

    Embranchement_Classe_Ordre_Famille_Genre_Espèce

    en :

    Embranchement_Classe_Ordre_Famille_Genre
    """

    if not taxon:
        return ""

    morceaux = str(taxon).strip().split("_")

    if len(morceaux) < 5:
        return str(taxon).strip()

    embranchement = morceaux[0]
    classe = morceaux[1]
    ordre = morceaux[2]
    famille = morceaux[3]
    genre = morceaux[4]

    return "_".join([
        embranchement,
        classe,
        ordre,
        famille,
        genre
    ])


def construire_taxon_adour(
    embranchement,
    classe,
    ordre,
    famille,
    genre
):
    """
    Construit le taxon Adour :

    Emb_Classis_Order_Familia_Genus
    """

    embranchement = champ_adour(
        embranchement,
        "Embranchement"
    )

    classe = champ_adour(
        classe,
        "Classis"
    )

    ordre = champ_adour(
        ordre,
        "Order"
    )

    famille = champ_adour(
        famille,
        "Familia"
    )

    genre = champ_adour(
        genre,
        "Genus"
    )

    if embranchement == "Embranchement":
        emb_court = "Emb"
    else:
        emb_court = embranchement

    if classe == "Classis":
        classe_court = "Classis"
    else:
        classe_court = classe

    if ordre == "Order":
        ordre_court = "Order"
    else:
        ordre_court = ordre

    return "_".join([
        emb_court,
        classe_court,
        ordre_court,
        famille,
        genre
    ])


# ============================================================
# NOUVELLE FONCTION DE COMPARAISON
# ============================================================


def est_generique(valeur):
    """
    Indique si une valeur est un terme taxonomique générique.
    """
    if valeur is None:
        return False

    valeur = str(valeur).strip().lower()  

    return valeur in {
        "emb",
        "embranchement",
        "classis",
        "order",
        "familia",
        "genus"
    }


def taxon_correspond(taxon1, taxon2):
    """
    Compare deux taxons en tenant compte des termes génériques.

    Règles :
    - Un générique suivi d'un autre générique reste générique.
    - Un générique suivi d'une information réelle est ignoré.
    - Le dernier 'Genus' reste toujours générique.
    - Un générique peut donc correspondre à n'importe quelle valeur
      au même niveau.
    - La comparaison est insensible à la casse.
    """

    if not taxon1 or not taxon2:
        return False

    champs1 = str(taxon1).strip().split("_")
    champs2 = str(taxon2).strip().split("_")

    # On travaille sur les 5 niveaux taxonomiques
    # Embranchement / Classe / Ordre / Famille / Genre
    if len(champs1) < 5 or len(champs2) < 5:
        return False

    champs1 = champs1[:5]
    champs2 = champs2[:5]

    for i in range(5):

        val1 = str(champs1[i]).strip().lower()  
        val2 = str(champs2[i]).strip().lower()  

        # ---------------------------------------------------------
        # Le dernier niveau : Genus
        # ---------------------------------------------------------
        if i == 4:

            # "Genus" est toujours générique
            if val1 == "genus" or val2 == "genus":
                continue

            if val1 != val2:
                return False

            continue

        # ---------------------------------------------------------
        # Les autres niveaux
        # ---------------------------------------------------------

        gen1 = est_generique(val1)
        gen2 = est_generique(val2)

        # Les deux sont génériques
        if gen1 and gen2:
            continue

        # Un seul est générique :
        # il est ignoré et l'autre valeur est acceptée
        if gen1 or gen2:
            continue

        # Les deux sont des valeurs réelles :
        # elles doivent être identiques
        if val1 != val2:
            return False

    return True


#--------------------------------------------------------------------------------------------------

# ============================================================
# INFORMATIONS
# ============================================================

print("********************************************************")
print()
if chemin.exists():
    print(f"Chemin utilisé : {chemin} : ") 
else:
    print(f"{chemin} : Chemin introuvable")
    exit()


if fichier_akkodis.exists():
    print(f"Fichier Akkodis : {fichier_akkodis.name}")
else:
    print(f"Fichier Akkodis : {fichier_akkodis.name} introuvable")
    exit()

if fichier_adour.exists():
    print(f"Fichier Adour : {fichier_adour.name}")
else:
    print(f"Fichier Adour : {fichier_adour.name} introuvable")
    exit()


# ============================================================
# CREATION DU FICHIER DE SORTIE
# ============================================================

wb_sortie = Workbook()

# Suppression de la feuille créée automatiquement
ws = wb_sortie.active
wb_sortie.remove(ws)

# ------------------------------------------------------------
# CREATION EXPLICITE DES 3 FEUILLES
# ------------------------------------------------------------

ws_akkodis = wb_sortie.create_sheet("Akkodis")
ws_adour = wb_sortie.create_sheet("Adour")
ws_synthese = wb_sortie.create_sheet("Synthèse")


# ============================================================
# FEUILLE 1 : AKKODIS
# ============================================================

# En-têtes

ws_akkodis["A1"] = "Taxon original"
ws_akkodis["B1"] = "NB"
ws_akkodis["C1"] = "T. référence"
ws_akkodis["D1"] = "Len"
ws_akkodis["E1"] = "Adour"
ws_akkodis["F1"] = "Fréquent"


# ============================================================
# LECTURE DES DONNEES AKKODIS
# ============================================================

wb_akkodis = load_workbook(
    fichier_akkodis,
    read_only=True,
    data_only=True
)

ws_akkodis_source = wb_akkodis.active

for ligne in ws_akkodis_source.iter_rows(
    min_row=2,
    values_only=True
):

    taxon = ligne[0]
    nombre = ligne[1]

    if taxon is None:
        continue

    taxon_modifie = construire_taxon_akkodis(taxon)

    nouvelle_ligne = ws_akkodis.max_row + 1

    # Colonne A : Taxon original
    ws_akkodis.cell(
        nouvelle_ligne,
        1,
        taxon
    )

    # Colonne B : NB
    ws_akkodis.cell(
        nouvelle_ligne,
        2,
        nombre
    )

    # Colonne C : Mis en Forme
    ws_akkodis.cell(
        nouvelle_ligne,
        3,
        taxon_modifie
    )

    # Colonne D : Len
    ws_akkodis.cell(
        nouvelle_ligne,
        4,
        len(taxon_modifie)
    )

wb_akkodis.close()


# ============================================================
# FEUILLE 2 : ADOUR
# ============================================================

# En-têtes

ws_adour["A1"] = "T. référence"
ws_adour["B1"] = "LEN"
ws_adour["C1"] = "Fréquent"
ws_adour["D1"] = "Akkodis"


# ============================================================
# OUVERTURE DATA_ADOUR
# ============================================================

wb_adour = load_workbook(
    fichier_adour,
    read_only=True,
    data_only=True
)

# La troisième feuille de Data_Adour
ws_adour_source = wb_adour.worksheets[2]


# ============================================================
# COLONNES UTILISEES DANS DATA_ADOUR
# ============================================================

col_regne = 1       # A = REGNE non utilisée
col_emb = 2         # B = EMBRANCHEMENT
col_classe = 3      # C = CLASSE
col_ordre = 4       # D = ORDRE
col_famille = 5     # E = FAMILLE
col_genre = 6       # F = GENRE
                       # G = Vide
col_frequent = 8    # H = Fréquent


# ============================================================
# LECTURE DES DONNEES ADOUR
# ============================================================

for ligne in ws_adour_source.iter_rows(
    min_row=2,
    values_only=False
):

    embranchement = ligne[col_emb - 1].value
    classe = ligne[col_classe - 1].value
    ordre = ligne[col_ordre - 1].value
    famille = ligne[col_famille - 1].value
    genre = ligne[col_genre - 1].value
    frequent = ligne[col_frequent - 1].value

    # Si toute la ligne taxonomique est vide,
    # on l'ignore
    if all(
        valeur is None
        for valeur in [
            embranchement,
            classe,
            ordre,
            famille,
            genre
        ]
    ):
        continue

    taxon_modifie = construire_taxon_adour(
        embranchement,
        classe,
        ordre,
        famille,
        genre
    )

    nouvelle_ligne = ws_adour.max_row + 1

    # Colonne A : Mis en Forme
    ws_adour.cell(
        nouvelle_ligne,
        1,
        taxon_modifie
    )

    # Colonne B : LEN
    ws_adour.cell(
        nouvelle_ligne,
        2,
        len(taxon_modifie)
    )

    # Colonne C : Fréquent
    ws_adour.cell(
        nouvelle_ligne,
        3,
        frequent
    )

wb_adour.close()


# ============================================================
# COMPARAISON AKKODIS / ADOUR
# ============================================================

print()
print("Comparaison des taxons Akkodis / Adour...")
print("Comparaison avec gestion des Termes génériques...")
print()


# ============================================================
# LISTE DES TAXONS ADOUR
# ============================================================

adour_taxons = []

for ligne in range(2, ws_adour.max_row + 1):

    taxon = ws_adour.cell(ligne, 1).value
    frequent = ws_adour.cell(ligne, 3).value

    if taxon is None:
        continue

    taxon = str(taxon).strip()

    if taxon:

        adour_taxons.append({
            "taxon": taxon,
            "frequent": frequent
        })


# ============================================================
# LISTE DES TAXONS AKKODIS
# ============================================================

akkodis_taxons = set()

for ligne in range(2, ws_akkodis.max_row + 1):

    taxon = ws_akkodis.cell(ligne, 3).value

    if taxon is None:
        continue

    taxon = str(taxon).strip()

    if taxon:
        akkodis_taxons.add(taxon)


# ============================================================
# PARCOURS AKKODIS
# ============================================================

# E = 1 si présent dans Adour
# F = Fréquent provenant de Adour

nb_akkodis_trouves = 0

for ligne in range(2, ws_akkodis.max_row + 1):

    taxon_akkodis = ws_akkodis.cell(ligne, 3).value

    if taxon_akkodis is None:
        continue

    taxon_akkodis = str(taxon_akkodis).strip()

    # Recherche d'un taxon correspondant dans Adour
    for element_adour in adour_taxons:

        taxon_adour = element_adour["taxon"]

        if taxon_correspond(taxon_akkodis, taxon_adour):

            # Colonne E : Adour
            ws_akkodis.cell(
                ligne,
                5,
                1
            )

            # Colonne F : Fréquent
            ws_akkodis.cell(
                ligne,
                6,
                element_adour["frequent"]
            )

            nb_akkodis_trouves += 1

            # On arrête dès qu'une correspondance est trouvée
            break


# ============================================================
# PARCOURS ADOUR
# ============================================================

# D = 1 si présent dans Akkodis

nb_adour_trouves = 0

for ligne in range(2, ws_adour.max_row + 1):

    taxon_adour = ws_adour.cell(ligne, 1).value

    if taxon_adour is None:
        continue

    taxon_adour = str(taxon_adour).strip()

    # Recherche d'une correspondance dans Akkodis
    for taxon_akkodis in akkodis_taxons:

        if taxon_correspond(taxon_adour, taxon_akkodis):

            # Colonne D : Akkodis
            ws_adour.cell(
                ligne,
                4,
                1
            )

            nb_adour_trouves += 1

            # Une correspondance suffit
            break


print(
    f"Taxons Akkodis trouvés dans Adour : "
    f"{nb_akkodis_trouves}"
)

print(
    f"Taxons Adour trouvés dans Akkodis : "
    f"{nb_adour_trouves}"
)


# ============================================================
# CREATION DE LA FEUILLE "SYNTHESE"
# ============================================================

# A = Colonne C de Akkodis, sans doublons
# B = Nombre de lignes regroupées
# C =  Présent dans Adour - Colonne E de Akkodis 
# D = Fréquent - Colonne F de Akkodis
# E = Somme de la colonne B de Akkodis

print("Création de la feuille Synthèse...")


# ============================================================
# EN-TÊTES
# ============================================================

ws_synthese["A1"] = "Taxons"
ws_synthese["B1"] = "Nb Reg"
ws_synthese["C1"] = ws_akkodis["E1"].value
ws_synthese["D1"] = ws_akkodis["F1"].value
ws_synthese["E1"] = "NB"


# ============================================================
# REGROUPEMENT DES DONNÉES AKKODIS
# ============================================================

regroupements = {}

for ligne in range(2, ws_akkodis.max_row + 1):

    # Clé de regroupement = colonne C
    taxon = ws_akkodis[f"C{ligne}"].value

    if taxon is None:
        continue

    taxon = str(taxon).strip()

    if not taxon:
        continue

    # Valeurs à récupérer
    valeur_e = ws_akkodis[f"E{ligne}"].value
    valeur_f = ws_akkodis[f"F{ligne}"].value
    valeur_b = ws_akkodis[f"B{ligne}"].value

    # Création du groupe
    if taxon not in regroupements:

        regroupements[taxon] = {
            "nombre": 0,
            "valeur_e": valeur_e,
            "valeur_f": valeur_f,
            "somme_b": 0
        }

    # Nombre de lignes regroupées
    regroupements[taxon]["nombre"] += 1

    # Somme de la colonne B
    if isinstance(valeur_b, (int, float)):
        regroupements[taxon]["somme_b"] += valeur_b


# ============================================================
# ÉCRITURE DE LA FEUILLE SYNTHÈSE
# ============================================================

ligne_synthese = 2

for taxon, valeurs in regroupements.items():

    # Colonne A : taxon unique provenant de C
    ws_synthese.cell(
        ligne_synthese,
        1,
        taxon
    )

    # Colonne B : nombre de regroupements
    ws_synthese.cell(
        ligne_synthese,
        2,
        valeurs["nombre"]
    )

    # Colonne C : informations provenant de E
    ws_synthese.cell(
        ligne_synthese,
        3,
        valeurs["valeur_e"]
    )

    # Colonne D : informations provenant de F
    ws_synthese.cell(
        ligne_synthese,
        4,
        valeurs["valeur_f"]
    )

    # Colonne E : somme de B
    ws_synthese.cell(
        ligne_synthese,
        5,
        valeurs["somme_b"]
    )

    ligne_synthese += 1


# ============================================================
# TOTAUX
# ============================================================

def calculer_totaux(feuille, col_somme, col_compte):

    total_data = feuille.max_row - 1

    tc = sum(
        feuille[f"{col_somme}{ligne}"].value or 0
        for ligne in range(2, feuille.max_row + 1)
        if isinstance(
            feuille[f"{col_somme}{ligne}"].value,
            (int, float)
        )
    )

    feuille[f"{col_compte}1"] = f"Total : {total_data}"

    feuille[f"{chr(ord(col_compte) + 1)}1"] = f"commun : {tc}"

    feuille[f"{chr(ord(col_compte) + 2)}1"] = (
        f"absent : {total_data - tc}"
    )


# Feuille 1 : Akkodis
calculer_totaux(ws_akkodis, 'E', 'H')

# Feuille 2 : Adour
calculer_totaux(ws_adour, 'D', 'H')

# Feuille 3 : Synthèse
calculer_totaux(ws_synthese, 'C', 'H')


# ============================================================
# MISE EN FORME
# ============================================================

def mise_en_forme_feuille(feuille, colonne):
    """
    Met en forme Première Ligne & deux colonnes consécutives.

    Paramètres :
        feuille  : feuille Excel openpyxl
        colonne  : première colonne à centrer, par exemple "C"
    """

    # Conversion de la lettre de colonne en numéro
    colonne = colonne.upper()
    col1 = column_index_from_string(colonne)
    col2 = col1 + 1

    # ---------------------------------------------------------
    # Première ligne : Gras + centrage horizontal et vertical
    # ---------------------------------------------------------
    for cell in feuille[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # ---------------------------------------------------------
    # Les deux colonnes : centrage sur toute la feuille
    # ---------------------------------------------------------
    for row in feuille.iter_rows():
        for cell in row:
            if cell.column in (col1, col2):
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

mise_en_forme_feuille(ws_akkodis, "E")
mise_en_forme_feuille(ws_adour, "C")
mise_en_forme_feuille(ws_synthese, "C")


# ============================================================
# SAUVEGARDE
# ============================================================

# Ouvrir directement sur la feuille "Synthèse"
wb_sortie.active = wb_sortie.index(ws_synthese)

wb_sortie.save(fichier_sortie)


# ============================================================
# VERIFICATION DES FEUILLES
# ============================================================

print()
print("==============================================")
print("Fichier créé avec succès")
print("==============================================")
print()

print("Fichier :", fichier_sortie.name)
print()

print("Feuilles créées :")

for numero, feuille in enumerate(
    wb_sortie.worksheets,
    start=1
):

    print(
        f"  {numero} - {feuille.title}"    # feuille.title
    )

print()

print("Structure de Synthèse :")
print("  A = Taxon")
print("  B = Nb de Regroupement")
print("  C = Présent dans Adour")
print("  D = Fréquent")
print("  E = Nombre d'élément")
print()
